"""Prepare explicitly reviewed private spreadsheet scopes; never activate old rates."""
from __future__ import annotations
import hashlib
import io
import json
import os
from decimal import Decimal
from pathlib import Path
import re
import zipfile

MAX_FILE_BYTES = 50 * 1024 * 1024


def prepare_entries(review: dict) -> list[dict]:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
    records = review.get('records')
    if not isinstance(records, list) or not 1 <= len(records) <= 50:
        raise ValueError('Select one to fifty reviewed source records')
    results = []
    for record in records:
        path = Path(record['source']).expanduser().resolve(strict=True)
        if path.suffix.lower() != '.xlsx' or path.stat().st_size > MAX_FILE_BYTES:
            raise ValueError('Use a bounded .xlsx source for this import preparation')
        data = path.read_bytes()
        # Offline, owner-selected history uses cached cell values only. External
        # workbook references are not followed (keep_links=False); active code,
        # embedded objects and archive expansion attacks remain disallowed.
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries)>5000 or sum(entry.file_size for entry in entries)>150*1024*1024:
                raise ValueError('Historical workbook expansion exceeds its limit')
            for entry in entries:
                name = entry.filename.lower()
                if '..' in Path(name).parts or name.startswith(('/', '\\')) or '\\' in name or any(token in name for token in ('vbaproject', 'macrosheets/', 'embeddings/', 'activex/')):
                    raise ValueError('Historical workbooks with active or embedded code are not supported')
                if entry.file_size>max(1,entry.compress_size)*1000:
                    raise ValueError('Historical workbook compression exceeds its limit')
            external_references = any('externallinks/' in entry.filename.lower() for entry in entries)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            locations = []
            def location(reference):
                if '!' in reference:
                    sheet, cell_range = reference.rsplit('!', 1)
                    sheet = sheet.strip("'")
                else:
                    sheet, cell_range = record.get('sourceSheet'), reference
                    if not sheet:
                        raise ValueError('Declare the reviewed source worksheet; do not infer it')
                if sheet not in workbook.sheetnames:
                    raise ValueError('A reviewed worksheet is missing from its original source')
                if not re.fullmatch(r'\$?[A-Z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Z]{1,3}\$?[1-9][0-9]*)?', cell_range):
                    raise ValueError('Use explicit reviewed spreadsheet cell ranges')
                low_col, low_row, high_col, high_row = range_boundaries(cell_range)
                if (high_col-low_col+1)*(high_row-low_row+1)>20000:
                    raise ValueError('Reviewed source range exceeds the extraction limit')
                cite = {'sheet': sheet, 'range': cell_range}
                if cite not in locations:
                    locations.append(cite)
                return workbook[sheet], cell_range, locations.index(cite)
            for reference in record.get('locations', []):
                location(reference)
            amounts = []
            for evidence in record.get('priceEvidence', []):
                sheet, reference, index = location(evidence['reference'])
                if ':' in reference:
                    raise ValueError('A monetary observation must name one exact source cell')
                role = evidence['role']
                if role not in ('selling', 'worker-cost', 'subtotal', 'total', 'alternative', 'unpriced', 'duplicate-total'):
                    raise ValueError('Classify each historical monetary observation')
                value = sheet[reference].value
                if role == 'unpriced':
                    if value not in (None, '', 0):
                        raise ValueError('A populated source cannot be labelled unpriced')
                    value_pence = None
                else:
                    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
                        raise ValueError('A historical amount requires a numeric cached value, never an inferred zero')
                    pence = Decimal(str(value))*100
                    if not pence.is_finite() or pence != pence.to_integral_value() or not 0 <= pence <= 1_000_000_000_000:
                        raise ValueError('The historical monetary observation is outside the supported precision')
                    value_pence = int(pence)
                amount = {'label': evidence['label'], 'role': role, 'valuePence': value_pence, 'locationIndex': index}
                if role == 'alternative':
                    if not evidence.get('alternativeGroup'):
                        raise ValueError('Alternatives need a group; they are never added together')
                    amount['alternativeGroup'] = evidence['alternativeGroup']
                amounts.append(amount)
            classification = {'historical-quotation': 'selling-quotation', 'historical-priced-schedule': 'selling-quotation', 'unpriced-tender': 'unpriced-schedule'}.get(record['classification'], record['classification'])
            if classification not in ('selling-quotation', 'worker-cost', 'unpriced-schedule', 'mixed-schedule'):
                raise ValueError('Classify the historical source before preparing it')
            warnings = list(record.get('warnings', []))
            if external_references:
                warnings.append('External workbook links were not followed or recalculated. Only reviewed cached historical cell values were read.')
            warnings.append('Historical observations only. Current owner-approved prices remain separate; no selling rates are imported.')
            results.append({'externalRef': record['id'], 'title': record.get('title', record['id']), 'classification': classification,
                'source': {'fileName': path.name, 'sha256': hashlib.sha256(data).hexdigest(), 'reviewedOn': record['reviewedOn'], 'locations': locations},
                'scopeNotes': record['scope'], 'quantityBasis': record['quantityBasis'], 'amounts': amounts, 'warnings': warnings})
        finally:
            workbook.close()
    return results


def write_private_export(review_path: Path, output_path: Path, public_repo: Path):
    output_path = output_path.expanduser().resolve()
    public_repo = public_repo.resolve()
    if output_path == public_repo or public_repo in output_path.parents:
        raise ValueError('Historical exports must remain outside the public website repository')
    if review_path.stat().st_size > 2_000_000:
        raise ValueError('Reviewed import manifest exceeds its limit')
    entries = prepare_entries(json.loads(review_path.read_text()))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation prevents overwriting an existing evidence artifact.
    descriptor = os.open(output_path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, 'w') as stream:
        json.dump({'entries': entries, 'currentRatesImported': False}, stream, indent=2)
        stream.write('\n')
    return {'records': len(entries), 'amountObservations': sum(len(entry['amounts']) for entry in entries), 'currentRatesImported': False}
