"""Bounded, offline document inspection. Source text is evidence, never a price instruction."""
from __future__ import annotations

import hashlib
import io
import re
import zipfile
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET

MAX_BYTES = 100 * 1024 * 1024
MAX_TEXT = 200_000
MAX_PAGES = 200
MAX_CELLS = 50_000
MIME = {
    '.pdf': 'application/pdf', '.png': 'image/png', '.jpg': 'image/jpeg',
    '.jpeg': 'image/jpeg', '.webp': 'image/webp', '.csv': 'text/csv',
    '.txt': 'text/plain', '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    '.dwg': 'application/octet-stream', '.dxf': 'application/octet-stream',
    '.ifc': 'application/octet-stream', '.rvt': 'application/octet-stream',
    '.mpp': 'application/octet-stream',
    '.heic': 'image/heic', '.heif': 'image/heif', '.tif': 'image/tiff', '.tiff': 'image/tiff',
    '.doc': 'application/msword', '.xls': 'application/vnd.ms-excel', '.ppt': 'application/vnd.ms-powerpoint',
    '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
}
NATIVE = {'.dwg', '.dxf', '.ifc', '.rvt', '.mpp', '.doc', '.xls', '.ppt', '.heic', '.heif', '.pptx'}


def safe_zip(data: bytes) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(io.BytesIO(data))
    members = archive.infolist()
    if len(members) > 5_000 or sum(x.file_size for x in members) > 150 * 1024 * 1024:
        raise ValueError('Office document expands beyond the processing limit')
    for item in members:
        p = PurePosixPath(item.filename)
        if p.is_absolute() or '..' in p.parts or '\\' in item.filename or item.flag_bits & 1:
            raise ValueError('Unsupported Office archive')
        if item.file_size > 10 * 1024 * 1024 and item.file_size > max(item.compress_size, 1) * 200:
            raise ValueError('Office archive compression limit exceeded')
        lower = item.filename.lower()
        if any(x in lower for x in ('vbaproject', 'macrosheets', 'externallinks', 'embeddings/')):
            raise ValueError('Embedded code or external workbook references are not supported')
    return archive


def inspect_pdf_objects(reader):
    """Inspect nested annotation/form actions too, with cycle and work bounds."""
    from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject
    visited, references = set(), set()
    stack = [(reader.trailer, 0)]
    count = 0
    while stack:
        node, depth = stack.pop()
        count += 1
        if count > 100_000 or depth > 100:
            raise ValueError('PDF object graph exceeds the processing limit')
        if isinstance(node, IndirectObject):
            key = (node.idnum, node.generation)
            if key in references:
                continue
            references.add(key)
            stack.append((node.get_object(), depth + 1))
            continue
        if isinstance(node, (DictionaryObject, ArrayObject)):
            if id(node) in visited:
                continue
            visited.add(id(node))
        if isinstance(node, DictionaryObject):
            if any(key in node for key in ('/OpenAction', '/AA', '/JavaScript', '/JS', '/EmbeddedFiles', '/EF', '/RichMediaContent', '/XFA')):
                raise ValueError('Active or embedded PDF content is not supported')
            if node.get('/S') in ('/JavaScript', '/Launch', '/SubmitForm', '/ImportData', '/GoToR', '/GoToE', '/Rendition', '/Movie', '/Sound'):
                raise ValueError('Active PDF action is not supported')
            if '/URI' in node and not re.match(r'^(?:https?://|mailto:)', str(node['/URI']), re.I):
                raise ValueError('Unsupported PDF link protocol')
            stack.extend((value, depth + 1) for value in node.values())
        elif isinstance(node, ArrayObject):
            stack.extend((value, depth + 1) for value in node)


def verify_document(data: bytes, file_name: str, expected_sha: str | None = None) -> dict:
    if not 0 < len(data) <= MAX_BYTES:
        raise ValueError('Document exceeds the file limit')
    ext = PurePosixPath(file_name).suffix.lower()
    if ext not in MIME:
        raise ValueError('Unsupported document format')
    digest = hashlib.sha256(data).hexdigest()
    if expected_sha and digest != expected_sha:
        raise ValueError('Document checksum mismatch')
    if ext == '.pdf':
        from pypdf import PdfReader
        if not data.startswith(b'%PDF-'):
            raise ValueError('File content does not match PDF')
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted or len(reader.pages) > MAX_PAGES:
            raise ValueError('Encrypted PDF or PDF page limit exceeded')
        inspect_pdf_objects(reader)
    elif ext in ('.png', '.jpg', '.jpeg', '.webp', '.tif', '.tiff'):
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = 30_000_000
        with Image.open(io.BytesIO(data)) as image:
            expected = {'.png': 'PNG', '.jpg': 'JPEG', '.jpeg': 'JPEG', '.webp': 'WEBP', '.tif': 'TIFF', '.tiff': 'TIFF'}[ext]
            if image.format != expected or image.width * image.height > 30_000_000 or getattr(image, 'n_frames', 1) != 1:
                raise ValueError('Image format or pixel limit is unsupported')
            image.verify()
    elif ext in ('.docx', '.xlsx', '.pptx'):
        with safe_zip(data) as archive:
            required = {'.docx':'word/document.xml', '.xlsx':'xl/workbook.xml', '.pptx':'ppt/presentation.xml'}[ext]
            if required not in archive.namelist():
                raise ValueError('File content does not match Office format')
    elif ext in ('.csv', '.txt'):
        text = data.decode('utf-8-sig')
        if '\x00' in text:
            raise ValueError('Text document contains binary content')
    elif ext in ('.doc', '.xls', '.ppt', '.mpp', '.rvt'):
        if not data.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
            raise ValueError('Source content does not match the compound document format')
    elif ext in ('.heic', '.heif'):
        if len(data) < 16 or data[4:8] != b'ftyp' or data[8:12] not in (b'heic', b'heix', b'hevc', b'hevx', b'mif1', b'msf1'):
            raise ValueError('Source content does not match HEIF')
    elif ext == '.dwg' and not re.match(rb'AC10\d\d', data[:6]):
        raise ValueError('Source content does not match DWG')
    elif ext == '.dxf' and b'SECTION' not in data[:1024] and not data.startswith(b'AutoCAD Binary DXF'):
        raise ValueError('Source content does not match DXF')
    elif ext == '.ifc' and not data.lstrip().startswith(b'ISO-10303-21;'):
        raise ValueError('Source content does not match IFC')
    return {'verified': True, 'mime': MIME[ext], 'sha256': digest, 'bytes': len(data), 'reviewRequired': ext in NATIVE}


def extract_evidence(data: bytes, file_name: str, document_id: str, *, ocr=None) -> dict:
    """Propose cited text and explicitly labelled measurements; never infer scale or geometry."""
    verify_document(data, file_name)
    ext = PurePosixPath(file_name).suffix.lower()
    evidence: list[dict] = []
    warnings: list[str] = []
    remaining = MAX_TEXT

    def add(text: str, **location):
        nonlocal remaining
        text = str(text).replace('\x00', '').strip()
        if not text or remaining <= 0 or len(evidence) >= 1996:
            if len(evidence) >= 1996:
                remaining = 0
            return
        text = text[:min(remaining, 4_000)]
        remaining -= len(text)
        source = {'documentId': document_id, **location}
        evidence.append({'id': f'e{len(evidence)+1}', 'kind': 'text', 'value': text, 'source': source, 'confirmed': False})
        # A printed label with an explicit unit is a proposal, not a confirmed wall.
        for match in re.finditer(r'\b(?:ceiling\s+height|wall\s+(?:height|width)|product\s+(?:code|reference))\s*[:=]\s*([^\n;]{1,100})', text, re.I):
            value = match.group(0)
            product = value.lower().startswith('product')
            if product or re.search(r'\d\s*(?:mm|cm|m|ft|in)\b', value):
                evidence.append({'id': f'e{len(evidence)+1}', 'kind': 'product_reference' if product else 'dimension', 'value': value, 'source': source, 'confirmed': False})

    if ext == '.pdf':
        from pypdf import PdfReader
        for number, page in enumerate(PdfReader(io.BytesIO(data)).pages, 1):
            text = page.extract_text() or ''
            if text.strip():
                add(text, page=number)
            else:
                warnings.append(f'Page {number}: no extractable text; provide a readable export or confirm measurements manually.')
            if remaining <= 0:
                break
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        cells = 0
        try:
            for sheet in formulas:
                for row in sheet:
                    for cell in row:
                        cells += 1
                        if cells > MAX_CELLS:
                            raise ValueError('Workbook cell limit exceeded')
                        if cell.value is None:
                            continue
                        value = values[sheet.title][cell.coordinate].value
                        if cell.data_type == 'f' and value is None:
                            warnings.append(f'{sheet.title}!{cell.coordinate}: formula has no cached value; never treated as zero.')
                            continue
                        add(str(value), sheet=sheet.title, cell=cell.coordinate)
                    if remaining <= 0:
                        break
        finally:
            formulas.close()
            values.close()
    elif ext == '.docx':
        with safe_zip(data) as archive:
            xml = archive.read('word/document.xml')
            if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml:
                raise ValueError('Unsupported XML declarations')
            root = ET.fromstring(xml)
            for index, paragraph in enumerate(root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'), 1):
                add(''.join(paragraph.itertext()), paragraph=index)
    elif ext in ('.txt', '.csv'):
        for number, line in enumerate(data.decode('utf-8-sig').splitlines(), 1):
            add(line, line=number)
            if remaining <= 0:
                break
    elif ext in NATIVE:
        warnings.append('Original source retained for specialist review. Upload a PDF export for document extraction.')
    elif ocr:
        for text in ocr(data):
            add(text, page=1)
    else:
        warnings.append('Image requires OCR or manual review. No dimensions have been inferred.')
    if remaining <= 0:
        warnings.append('Text limit reached; review the remaining source manually.')
    return {'evidence': evidence, 'warnings': warnings, 'requiresConfirmation': True, 'extractorVersion': 'designer-documents-1'}
